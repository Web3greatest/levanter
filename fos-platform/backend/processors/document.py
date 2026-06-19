"""
Document Processor — handles PDF, DOCX, PPTX, CSV, Excel, TXT, and images.
Chunks text, extracts data, and prepares content for the memory/RAG system.
"""
import os
import io
import uuid
import asyncio
from pathlib import Path
from typing import Optional

import aiofiles
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select

from database import Document
from config import settings


def _chunk_text(text: str, chunk_size: int = None, overlap: int = None) -> list[str]:
    """Split text into overlapping chunks."""
    chunk_size = chunk_size or settings.CHUNK_SIZE
    overlap = overlap or settings.CHUNK_OVERLAP

    if len(text) <= chunk_size:
        return [text]

    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        # Try to break at sentence boundary
        if end < len(text):
            for sep in ['. ', '.\n', '\n\n', '\n', ' ']:
                idx = text.rfind(sep, start, end)
                if idx > start + overlap:
                    end = idx + len(sep)
                    break
        chunks.append(text[start:end].strip())
        start = end - overlap

    return [c for c in chunks if c]


def _extract_pdf(file_path: str) -> tuple[str, int]:
    """Extract text from PDF. Returns (text, page_count)."""
    try:
        import PyPDF2
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            pages = []
            for page in reader.pages:
                pages.append(page.extract_text() or "")
            return "\n\n".join(pages), len(reader.pages)
    except Exception:
        # Fallback to pdfminer
        try:
            from pdfminer.high_level import extract_text as pdfminer_extract
            text = pdfminer_extract(file_path)
            return text, 0
        except Exception as e:
            return f"[PDF extraction error: {str(e)}]", 0


def _extract_docx(file_path: str) -> str:
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also extract tables
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        return "\n\n".join(paragraphs)
    except Exception as e:
        return f"[DOCX extraction error: {str(e)}]"


def _extract_pptx(file_path: str) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            if texts:
                slides.append(f"Slide {i}:\n" + "\n".join(texts))
        return "\n\n".join(slides)
    except Exception as e:
        return f"[PPTX extraction error: {str(e)}]"


def _extract_csv(file_path: str) -> str:
    try:
        import pandas as pd
        df = pd.read_csv(file_path, nrows=1000)  # limit rows
        return f"Columns: {', '.join(df.columns)}\n\n" + df.to_string(max_rows=100)
    except Exception as e:
        return f"[CSV extraction error: {str(e)}]"


def _extract_excel(file_path: str) -> str:
    try:
        import pandas as pd
        xl = pd.ExcelFile(file_path)
        parts = []
        for sheet in xl.sheet_names[:5]:  # max 5 sheets
            df = xl.parse(sheet, nrows=500)
            parts.append(f"Sheet: {sheet}\nColumns: {', '.join(str(c) for c in df.columns)}\n{df.to_string(max_rows=50)}")
        return "\n\n---\n\n".join(parts)
    except Exception as e:
        return f"[Excel extraction error: {str(e)}]"


def _extract_image_ocr(file_path: str) -> str:
    try:
        import pytesseract
        from PIL import Image
        img = Image.open(file_path)
        text = pytesseract.image_to_string(img)
        return text.strip() or "[No text found in image]"
    except Exception as e:
        return f"[OCR error: {str(e)}]"


class DocumentProcessor:
    async def process_file(self, file_path: str, filename: str) -> dict:
        """Process a file and return extracted content."""
        ext = Path(filename).suffix.lower().lstrip('.')
        size = os.path.getsize(file_path)

        # Run extraction in thread pool (CPU-bound)
        loop = asyncio.get_event_loop()

        if ext == 'pdf':
            text, pages = await loop.run_in_executor(None, _extract_pdf, file_path)
            page_count = pages
        elif ext in ('doc', 'docx'):
            text = await loop.run_in_executor(None, _extract_docx, file_path)
            page_count = None
        elif ext in ('ppt', 'pptx'):
            text = await loop.run_in_executor(None, _extract_pptx, file_path)
            page_count = None
        elif ext == 'csv':
            text = await loop.run_in_executor(None, _extract_csv, file_path)
            page_count = None
        elif ext in ('xlsx', 'xls'):
            text = await loop.run_in_executor(None, _extract_excel, file_path)
            page_count = None
        elif ext in ('jpg', 'jpeg', 'png', 'gif', 'bmp', 'webp'):
            text = await loop.run_in_executor(None, _extract_image_ocr, file_path)
            page_count = 1
        elif ext == 'txt':
            async with aiofiles.open(file_path, 'r', errors='replace') as f:
                text = await f.read()
            page_count = None
        else:
            text = f"Unsupported file type: {ext}"
            page_count = None

        chunks = _chunk_text(text)

        return {
            "text": text,
            "chunks": chunks,
            "page_count": page_count,
            "size_bytes": size,
            "file_type": ext,
        }

    async def save_document(
        self,
        db: AsyncSession,
        user_id: str,
        filename: str,
        file_content: bytes,
        ai_router=None,
    ) -> Document:
        """Save uploaded file, process it, and store in DB."""
        doc_id = str(uuid.uuid4())
        file_path = os.path.join(settings.UPLOAD_DIR, f"{doc_id}_{filename}")

        async with aiofiles.open(file_path, 'wb') as f:
            await f.write(file_content)

        processed = await self.process_file(file_path, filename)

        summary = ""
        if ai_router and processed["text"] and len(processed["text"]) > 100:
            try:
                text_sample = processed["text"][:6000]
                messages = [
                    {"role": "system", "content": "Summarize this document in 3-5 sentences, capturing the key topics, data, and insights."},
                    {"role": "user", "content": text_sample}
                ]
                from ai.router import TaskComplexity
                summary, _ = await ai_router.complete(messages, complexity=TaskComplexity.FAST)
            except Exception:
                summary = processed["text"][:200] + "..."

        doc = Document(
            id=doc_id,
            user_id=user_id,
            filename=filename,
            file_type=processed["file_type"],
            file_path=file_path,
            size_bytes=processed["size_bytes"],
            page_count=processed["page_count"],
            raw_text=processed["text"][:50000],  # store first 50K chars
            summary=summary,
            chunks=processed["chunks"][:200],  # max 200 chunks
        )

        db.add(doc)
        await db.commit()
        await db.refresh(doc)
        return doc

    async def get_document(self, db: AsyncSession, doc_id: str, user_id: str) -> Optional[Document]:
        result = await db.execute(
            select(Document).where(Document.id == doc_id, Document.user_id == user_id)
        )
        return result.scalar_one_or_none()

    async def list_documents(self, db: AsyncSession, user_id: str) -> list[Document]:
        result = await db.execute(
            select(Document).where(Document.user_id == user_id)
            .order_by(Document.uploaded_at.desc())
        )
        return list(result.scalars())

    async def search_documents(self, db: AsyncSession, user_id: str, query: str, k: int = 3) -> str:
        """Simple keyword search across document chunks."""
        docs = await self.list_documents(db, user_id)
        if not docs:
            return ""

        query_words = set(query.lower().split())
        scored = []

        for doc in docs:
            for chunk in (doc.chunks or [])[:50]:  # check first 50 chunks
                chunk_lower = chunk.lower()
                hits = sum(1 for w in query_words if w in chunk_lower)
                if hits > 0:
                    scored.append((chunk, hits, doc.filename))

        scored.sort(key=lambda x: x[1], reverse=True)
        top = scored[:k]

        if not top:
            return ""

        parts = [f"From '{fn}':\n{chunk}" for chunk, _, fn in top]
        return "\n\n".join(parts)

    async def answer_from_document(
        self,
        db: AsyncSession,
        doc_id: str,
        user_id: str,
        question: str,
        ai_router,
    ) -> str:
        """Answer a question using a specific document's content."""
        doc = await self.get_document(db, doc_id, user_id)
        if not doc:
            return "Document not found."

        # Find most relevant chunks
        query_words = set(question.lower().split())
        scored_chunks = []
        for chunk in (doc.chunks or []):
            hits = sum(1 for w in query_words if w in chunk.lower())
            scored_chunks.append((chunk, hits))
        scored_chunks.sort(key=lambda x: x[1], reverse=True)
        relevant = "\n\n".join(c for c, _ in scored_chunks[:8])

        messages = [
            {
                "role": "system",
                "content": f"You are analyzing the document '{doc.filename}'. Answer questions based only on the provided content. If the answer is not in the content, say so."
            },
            {
                "role": "user",
                "content": f"Document content:\n{relevant}\n\nQuestion: {question}"
            }
        ]

        from ai.router import TaskComplexity
        answer, _ = await ai_router.complete(messages, complexity=TaskComplexity.SMART)
        return answer


document_processor = DocumentProcessor()
